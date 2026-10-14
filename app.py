import os
import json
import re
from io import BytesIO

import requests
import streamlit as st

st.set_page_config(
    page_title="ATS Resume Copilot",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
.block-container {max-width: 1180px; padding-top: 2rem; padding-bottom: 3rem;}
.hero {padding: 1.5rem 1.7rem; border-radius: 18px; background: linear-gradient(120deg,#10223d,#164e63); color: white; margin-bottom: 1.2rem;}
.hero h1 {font-size: 2.25rem; margin: 0 0 .35rem 0;}
.hero p {font-size: 1.02rem; opacity: .92; margin: 0;}
.small-note {color: #6b7280; font-size: .88rem;}
.step-card {border: 1px solid rgba(128,128,128,.25); border-radius: 14px; padding: 1rem; height: 100%;}
div[data-testid="stMetric"] {border: 1px solid rgba(128,128,128,.2); padding: 12px; border-radius: 12px;}
</style>
<div class="hero">
  <h1>🎯 ATS Resume Copilot</h1>
  <p>Tailor your resume to a real job description — analyze, score, improve, and export.</p>
</div>
""", unsafe_allow_html=True)

with st.sidebar:
    st.header("⚙️ AI settings")
    provider = st.selectbox("AI provider", ["Ollama (local)", "OpenAI-compatible API"], index=0)
    if provider == "Ollama (local)":
        endpoint = st.text_input("Ollama endpoint", "http://localhost:11434")
        model = st.text_input("Model name", "llama3.1")
        api_key = ""
        st.caption("Ollama must be running on your computer. Pull a model first, e.g. `ollama pull llama3.1`.")
    else:
        endpoint = st.text_input("Chat completions endpoint", "https://api.openai.com/v1/chat/completions")
        model = st.text_input("Model name", "gpt-4o-mini")
        api_key = st.text_input("API key", type="password")
        st.caption("Your provider may charge for API usage. The app does not store your key.")
    st.divider()
    st.subheader("Privacy")
    st.caption("Resume text is sent to the AI provider you configure. With local Ollama, processing can stay on your computer.")

def extract_text(uploaded):
    name = uploaded.name.lower()
    data = uploaded.getvalue()
    if name.endswith(".pdf"):
        try:
            import fitz
            doc = fitz.open(stream=data, filetype="pdf")
            return "\n".join(page.get_text("text") for page in doc).strip()
        except Exception as e:
            raise ValueError("Could not read this PDF. Try a text-based PDF or upload DOCX/TXT.") from e
    if name.endswith(".docx"):
        try:
            from docx import Document
            doc = Document(BytesIO(data))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts).strip()
        except Exception as e:
            raise ValueError("Could not read this DOCX file.") from e
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="replace").strip()
    raise ValueError("Upload a PDF, DOCX, or TXT file.")

SYSTEM_PROMPT = """You are an expert technical recruiter and ATS resume specialist. Tailor a candidate's resume to the supplied job description.
Rules:
- Never invent employers, dates, titles, qualifications, skills, metrics, certifications, project features, or achievements.
- Use only evidence in the original resume. If a claim is uncertain, flag it for verification instead of treating it as fact.
- Distinguish required skills from preferred skills.
- The match score is an illustrative rubric-based estimate, NOT an actual ATS score and NOT a callback prediction.
- Do not guarantee interviews or shortlisting.
- Keep language concise, clear, and natural. Avoid keyword stuffing.
- Do not change factual job titles or employment dates.
- Keep the resume ATS-friendly: single-column, standard headings, no tables in the resume body.
Return valid JSON only. Do not use markdown fences."""

USER_TEMPLATE = """Perform these five steps automatically using the resume and job description below.

STEP 1 — Analyze the job description:
List top 10 requirements, split required vs preferred, key tools/skills, responsibilities, and eligibility criteria.

STEP 2 — Review current resume:
Identify strengths, weak/generic bullets, missing evidence, formatting issues, and specific risks.

STEP 3 — Score match transparently:
Use these weights: required skills 30, relevant experience/projects 25, responsibilities/business relevance 20, education/eligibility 15, preferred qualifications 10. Return each category score and explain evidence. Do not treat preferred skills as mandatory.

STEP 4 — Tailor resume:
Return a complete resume ready to copy. Prioritize relevant existing evidence and use action verbs. Keep skills truthful. Include contact links only if present in source. Mark unclear claims for verification. Recommend a length, but do not force tiny text.

STEP 5 — Final gap and application plan:
List missing keywords/skills (distinguish absent from not demonstrated), rejection risks with fixes, final ATS checklist, and five likely interview questions based on the resume/JD.

Return JSON with exactly these top-level keys:
"job_analysis": {"top_requirements": [], "required_skills": [], "preferred_skills": [], "responsibilities": [], "eligibility": []},
"resume_review": {"strengths": [], "weaknesses": [], "formatting_issues": [], "claims_to_verify": []},
"match_score": {"total": integer, "out_of": 100, "categories": [{"name": string, "score": number, "max": number, "reason": string}], "explanation": string},
"missing_keywords": [{"keyword": string, "status": "missing" or "not_demonstrated", "importance": "high" or "medium" or "low", "note": string}],
"tailored_resume": "complete plain-text resume",
"rejection_risks": [{"risk": string, "why_it_matters": string, "fix": string}],
"checklist": [],
"interview_questions": [],
"recommended_filename": string

JOB DESCRIPTION:
{jd}

ORIGINAL RESUME:
{resume}
"""

def call_ai(provider, endpoint, model, api_key, resume, jd):
    prompt = USER_TEMPLATE.replace("{jd}", jd).replace("{resume}", resume)
    if provider == "Ollama (local)":
        url = endpoint.rstrip("/") + "/api/chat"
        payload = {
            "model": model,
            "messages": [
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": prompt},
            ],
            "stream": False,
            "format": "json",
            "options": {"temperature": 0.2},
        }
        response = requests.post(url, json=payload, timeout=240)
        response.raise_for_status()
        content = response.json()["message"]["content"]
    else:
        if not api_key:
            raise ValueError("Enter an API key in the sidebar, or use local Ollama.")
        headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}
        payload = {
            "model": model,
            "messages": [
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user", "content": prompt},
            ],
            "temperature": 0.2,
            "response_format": {"type": "json_object"},
        }
        response = requests.post(endpoint, headers=headers, json=payload, timeout=240)
        response.raise_for_status()
        content = response.json()["choices"][0]["message"]["content"]
    # Strip accidental code fences from providers that don't honor JSON mode.
    content = re.sub(r"^\s*```(?:json)?\s*|\s*```\s*$", "", content, flags=re.I)
    return json.loads(content)

st.subheader("1 · Add your resume and target job")
left, right = st.columns([0.9, 1.1], gap="large")
with left:
    uploaded = st.file_uploader("Upload current resume", type=["pdf", "docx", "txt"])
    resume_text_manual = st.text_area("Or paste resume text", height=180, placeholder="Paste resume text here if you don't want to upload a file.")
with right:
    target_role = st.text_input("Target role (optional)", placeholder="e.g. Junior Data Scientist")
    location = st.text_input("Preferred location (optional)", placeholder="e.g. Hyderabad / Remote India")
    jd = st.text_area("Paste the full job description *", height=260, placeholder="Paste responsibilities, required skills, qualifications, and experience requirements.")

resume_text = resume_text_manual.strip()
if uploaded is not None:
    try:
        extracted = extract_text(uploaded)
        if extracted:
            resume_text = extracted
            st.success(f"Resume loaded: {uploaded.name} · {len(extracted.split())} words extracted")
        else:
            st.warning("No selectable text found. This may be a scanned PDF; paste the text or OCR it first.")
    except Exception as exc:
        st.error(str(exc))

if target_role or location:
    st.caption("Context: " + " · ".join(x for x in [target_role, location] if x))

run = st.button("🚀 Run all 5 steps", type="primary", use_container_width=True)
if run:
    if not resume_text:
        st.error("Upload a resume or paste its text.")
    elif not jd.strip():
        st.error("Paste the target job description.")
    else:
        with st.spinner("Analyzing the job, checking the resume, scoring the match, and drafting tailored output…"):
            try:
                enriched_jd = jd.strip()
                if target_role.strip():
                    enriched_jd = f"Target role: {target_role.strip()}\nPreferred location: {location.strip()}\n\n" + enriched_jd
                result = call_ai(provider, endpoint, model, api_key, resume_text, enriched_jd)
                st.session_state["result"] = result
                st.session_state["source_resume"] = resume_text
                st.session_state["source_jd"] = enriched_jd
            except requests.exceptions.ConnectionError:
                st.error("Could not connect to the AI endpoint. For Ollama, make sure Ollama is running and the model is downloaded. For an API, verify the endpoint.")
            except requests.exceptions.HTTPError as exc:
                detail = ""
                try:
                    detail = exc.response.text[:1000]
                except Exception:
                    pass
                st.error(f"AI provider returned an HTTP error: {exc}. {detail}")
            except Exception as exc:
                st.error(f"Could not complete the analysis: {exc}")

result = st.session_state.get("result")
if result:
    st.divider()
    st.subheader("2 · Your results")
    score = result.get("match_score", {})
    c1, c2, c3 = st.columns(3)
    c1.metric("Resume-to-job match", f'{score.get("total", "—")}/{score.get("out_of", 100)}')
    c2.metric("Top requirements", len(result.get("job_analysis", {}).get("top_requirements", [])))
    c3.metric("Items to verify", len(result.get("resume_review", {}).get("claims_to_verify", [])))
    st.caption("Match score is an illustrative rubric-based estimate, not an actual ATS score or callback prediction.")

    tabs = st.tabs(["Job analysis", "Resume review", "Match score", "Missing keywords", "Tailored resume", "Risks & checklist", "Interview prep"])
    with tabs[0]:
        ja = result.get("job_analysis", {})
        for title, key in [("Top requirements", "top_requirements"), ("Required skills", "required_skills"), ("Preferred skills", "preferred_skills"), ("Responsibilities", "responsibilities"), ("Eligibility", "eligibility")]:
            with st.expander(title, expanded=(key == "top_requirements")):
                vals = ja.get(key, [])
                if vals:
                    for item in vals:
                        st.markdown(f"- {item}")
                else:
                    st.write("No items returned.")
    with tabs[1]:
        rr = result.get("resume_review", {})
        for title, key in [("Strengths", "strengths"), ("Weaknesses", "weaknesses"), ("Formatting issues", "formatting_issues"), ("Claims to verify", "claims_to_verify")]:
            st.markdown(f"**{title}**")
            vals = rr.get(key, [])
            if vals:
                for item in vals:
                    st.markdown(f"- {item}")
            else:
                st.write("None identified.")
    with tabs[2]:
        cats = score.get("categories", [])
        if cats:
            st.dataframe([{"Category": x.get("name"), "Score": x.get("score"), "Maximum": x.get("max"), "Evidence": x.get("reason")} for x in cats], use_container_width=True, hide_index=True)
        st.write(score.get("explanation", ""))
    with tabs[3]:
        kws = result.get("missing_keywords", [])
        if kws:
            st.dataframe([{"Keyword": x.get("keyword"), "Status": x.get("status"), "Importance": x.get("importance"), "Note": x.get("note")} for x in kws], use_container_width=True, hide_index=True)
        else:
            st.success("No missing keywords were identified by the analysis.")
        st.caption("Do not add a missing skill unless you genuinely have it.")
    with tabs[4]:
        tailored = result.get("tailored_resume", "")
        st.text_area("Edit or copy your tailored resume", value=tailored, height=520, key="tailored_editor")
        st.download_button("⬇️ Download tailored resume (.txt)", data=st.session_state.get("tailored_editor", tailored), file_name=result.get("recommended_filename", "tailored_resume.txt"), mime="text/plain", use_container_width=True)
        st.info("Review all dates, links, skills, and metrics before applying. This export is editable text; format it in Word/Google Docs and export as PDF.")
    with tabs[5]:
        risks = result.get("rejection_risks", [])
        for r in risks:
            with st.expander(r.get("risk", "Potential risk")):
                st.write("**Why it matters:** " + r.get("why_it_matters", ""))
                st.write("**Recommended fix:** " + r.get("fix", ""))
        st.markdown("**Final ATS checklist**")
        for item in result.get("checklist", []):
            st.markdown(f"- [ ] {item}")
    with tabs[6]:
        for i, q in enumerate(result.get("interview_questions", []), start=1):
            st.markdown(f"**{i}.** {q}")

    export = json.dumps(result, ensure_ascii=False, indent=2)
    st.download_button("⬇️ Download complete analysis (.json)", data=export, file_name="ats_resume_analysis.json", mime="application/json")
    st.caption("Privacy reminder: review the AI provider's data policy before uploading sensitive information. Remove personal identifiers if not needed.")

st.divider()
st.markdown("""
<div class="small-note">
ATS Resume Copilot is a resume-assistance tool, not an employer ATS. It cannot guarantee callbacks or shortlisting.
Always verify generated content before applying. Keep your resume factual; never add a keyword or achievement you cannot support.
</div>
""", unsafe_allow_html=True)
